import io
from typing import Dict, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="1E293B")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="64748B")
DATA_FONT = Font(name="Calibri", size=10)
TOTAL_FONT = Font(name="Calibri", bold=True, size=11, color="059669")
BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def export_boq_to_excel(boq_data: Dict, project_name: str = "BOQ Report") -> bytes:
    wb = Workbook()

    # === Sheet 1: BOQ Summary ===
    ws = wb.active
    ws.title = "BOQ Summary"
    ws.sheet_properties.tabColor = "059669"

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 18

    # Title
    ws.merge_cells("A1:G1")
    cell = ws["A1"]
    cell.value = f"Bill of Quantities — {project_name}"
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal="right")

    ws.merge_cells("A2:G2")
    ws["A2"].value = f"Total Elements: {boq_data.get('elements_count', 0)} | Classified: {boq_data.get('classified_count', 0)} | Unclassified: {boq_data.get('unclassified_count', 0)}"
    ws["A2"].font = SUBTITLE_FONT

    # Summary section
    ws["A4"].value = "COST SUMMARY"
    ws["A4"].font = Font(name="Calibri", bold=True, size=12, color="059669")

    summary = boq_data.get("summary", {})
    summary_rows = [
        ("Base Cost", summary.get("total_base_cost", 0)),
        ("Overhead (10%)", summary.get("overhead", 0)),
        ("Profit (15%)", summary.get("profit", 0)),
        ("Contingency (5%)", summary.get("contingency", 0)),
        ("VAT (15%)", summary.get("vat", 0)),
        ("TOTAL", summary.get("total", 0)),
    ]

    for i, (label, value) in enumerate(summary_rows):
        row = 5 + i
        ws[f"B{row}"].value = label
        ws[f"B{row}"].font = TOTAL_FONT if label == "TOTAL" else DATA_FONT
        ws[f"G{row}"].value = value
        ws[f"G{row}"].font = TOTAL_FONT if label == "TOTAL" else DATA_FONT
        ws[f"G{row}"].number_format = '#,##0.00'

    # BOQ Table Header
    table_start = 5 + len(summary_rows) + 2
    ws.merge_cells(f"A{table_start}:G{table_start}")
    ws[f"A{table_start}"].value = "DETAILED BOQ"
    ws[f"A{table_start}"].font = Font(name="Calibri", bold=True, size=12, color="059669")

    header_row = table_start + 1
    headers = ["#", "Element Type", "Description", "Unit", "Quantity", "Unit Price", "Total"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER

    by_type = boq_data.get("by_type", {})
    row = header_row + 1
    idx = 1
    for et_key, et_data in by_type.items():
        if isinstance(et_data, dict):
            ws.cell(row=row, column=1, value=idx).font = DATA_FONT
            ws.cell(row=row, column=2, value=et_key).font = DATA_FONT
            ws.cell(row=row, column=3, value=et_data.get("element_type", "")).font = DATA_FONT
            ws.cell(row=row, column=4, value=et_data.get("unit", "م2")).font = DATA_FONT
            ws.cell(row=row, column=5, value=et_data.get("total_quantity", 0)).font = DATA_FONT
            ws.cell(row=row, column=5).number_format = '#,##0.00'
            ws.cell(row=row, column=6).value = ""
            ws.cell(row=row, column=7, value=et_data.get("base_cost", 0)).font = DATA_FONT
            ws.cell(row=row, column=7).number_format = '#,##0.00'
            for c in range(1, 8):
                ws.cell(row=row, column=c).border = BORDER
            row += 1
            idx += 1

    # === Sheet 2: By Drawing ===
    ws2 = wb.create_sheet("By Drawing")
    ws2.column_dimensions["A"].width = 15
    ws2.column_dimensions["B"].width = 15
    ws2.column_dimensions["C"].width = 15
    ws2.column_dimensions["D"].width = 15
    ws2.column_dimensions["E"].width = 18

    ws2["A1"].value = "BOQ BY DRAWING"
    ws2["A1"].font = TITLE_FONT

    headers2 = ["Drawing ID", "Elements", "Classified", "Total Quantity", "Base Cost"]
    for col, h in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER

    row2 = 4
    for d in boq_data.get("by_drawing", []):
        ws2.cell(row=row2, column=1, value=d.get("drawing_id", "")).font = DATA_FONT
        ws2.cell(row=row2, column=2, value=d.get("elements_count", 0)).font = DATA_FONT
        ws2.cell(row=row2, column=3, value=d.get("classified_count", 0)).font = DATA_FONT
        ws2.cell(row=row2, column=4, value=d.get("total_quantity", 0)).font = DATA_FONT
        ws2.cell(row=row2, column=4).number_format = '#,##0.00'
        ws2.cell(row=row2, column=5, value=d.get("base_cost", 0)).font = DATA_FONT
        ws2.cell(row=row2, column=5).number_format = '#,##0.00'
        for c in range(1, 6):
            ws2.cell(row=row2, column=c).border = BORDER
        row2 += 1

    # Output
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
