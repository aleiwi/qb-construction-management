"""Excel/PDF export for the quick villa estimator (legacy had onPdf/onExcel)."""
import io
from datetime import datetime
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor, white
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from app.services.arabic_pdf import register_arabic_font, shape_ar

DARK = HexColor("#1E293B")
EMERALD = HexColor("#059669")
LIGHT = HexColor("#F8FAFC")
BORDER = HexColor("#CBD5E1")

HEADERS = ["البند", "خرسانة (م³)", "حديد (طن)", "بلوك", "المسطح (م²)", "الإجمالي (ر.س)"]


def export_estimate_excel(items: List[Dict], totals: Dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "التقدير"
    ws.sheet_view.rightToLeft = True
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1E293B")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    ws.append(HEADERS)
    for c in ws[1]:
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border
    for it in items:
        row = [it["name"], it["concrete_m3"], it["steel_tons"], it["blocks"],
               it["area_m2"], it["total_cost"]]
        ws.append(row)
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = right
    ws.append([])
    for label, key in [("إجمالي المواد", "materials_cost"), ("المصنعية", "contractor_cost"),
                       ("قبل الربح", "total_before_profit"), ("الربح", "profit"),
                       ("السعر النهائي", "final_price"), ("التكلفة/م²", "cost_per_m2")]:
        ws.append([label, totals.get(key, 0)])
        ws[ws.max_row][0].font = Font(bold=True)
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = right
    for col in range(1, 7):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_estimate_pdf(items: List[Dict], totals: Dict, title: str = "تقدير سريع — فلل سكنية") -> bytes:
    font = register_arabic_font()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ArTitle", fontName=font, fontSize=20, alignment=1, spaceAfter=6))
    styles.add(ParagraphStyle("ArBody", fontName=font, fontSize=10, alignment=2, spaceAfter=4))
    styles.add(ParagraphStyle("ArCell", fontName=font, fontSize=9, alignment=1))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm,
                            leftMargin=2 * cm, rightMargin=2 * cm)
    els = [Paragraph(shape_ar(title), styles["ArTitle"]),
           Paragraph(shape_ar(f"التاريخ: {datetime.utcnow().strftime('%Y-%m-%d')}"), styles["ArBody"]),
           Spacer(1, 12)]
    data = [[Paragraph(shape_ar(h), styles["ArCell"]) for h in HEADERS]]
    for it in items:
        data.append([
            Paragraph(shape_ar(it["name"]), styles["ArCell"]),
            f"{it['concrete_m3']:,.2f}", f"{it['steel_tons']:,.3f}",
            f"{it['blocks']:,.2f}", f"{it['area_m2']:,.2f}", f"{it['total_cost']:,.2f}",
        ])
    t = Table(data, colWidths=[120, 60, 60, 60, 60, 80], repeatRows=1)
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), DARK),
        ("TEXTCOLOR", (0, 0), (-1, 0), white),
        ("GRID", (0, 0), (-1, -1), 0.5, BORDER),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [white, LIGHT]),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    els.append(t)
    els.append(Spacer(1, 12))
    for label, key in [("إجمالي المواد", "materials_cost"), ("السعر النهائي", "final_price"),
                       ("التكلفة/م²", "cost_per_m2")]:
        els.append(Paragraph(
            shape_ar(f"{label}: {totals.get(key, 0):,.2f} ر.س"), styles["ArBody"]))
    doc.build(els)
    buf.seek(0)
    return buf.getvalue()
